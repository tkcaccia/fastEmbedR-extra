#!/usr/bin/env python3

import sys
from copy import deepcopy

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_section_break(anchor, template, landscape):
    paragraph = OxmlElement("w:p")
    properties = OxmlElement("w:pPr")
    section = deepcopy(template)
    size = section.find(qn("w:pgSz"))
    margins = section.find(qn("w:pgMar"))
    if size is None:
        size = OxmlElement("w:pgSz")
        section.insert(0, size)
    if margins is None:
        margins = OxmlElement("w:pgMar")
        section.append(margins)
    if landscape:
        size.set(qn("w:w"), str(Inches(11.69).twips))
        size.set(qn("w:h"), str(Inches(8.27).twips))
        size.set(qn("w:orient"), "landscape")
        margins.set(qn("w:left"), str(Inches(0.55).twips))
        margins.set(qn("w:right"), str(Inches(0.55).twips))
    properties.append(section)
    paragraph.append(properties)
    anchor.addprevious(paragraph)


def format_wide_table(table):
    widths = [Inches(2.2)] + [Inches(0.75)] * 11
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = width
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(7)


def format_capability_table(table):
    widths = [Inches(value) for value in
              (1.40, 0.85, 0.65, 0.49, 0.49, 1.05, 0.77, 0.77)]
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = width
    for row in table.rows:
        for index, (cell, width) in enumerate(zip(row.cells, widths)):
            cell.width = width
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                if index >= 2:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(8)


def format_workflow_table(table):
    widths = [Inches(value) for value in
              (1.10, 1.00, 0.95, 1.10, 1.25, 1.10)]
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = width
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(8)


def format_author_superscripts(document):
    for paragraph in document.paragraphs:
        if paragraph.style.name != "Author":
            continue
        for math in paragraph._element.xpath("./m:oMath"):
            markers = [node for sup in math.iter()
                       if sup.tag == qn("m:sup")
                       for node in sup.iter()
                       if node.tag == qn("m:t")]
            if not markers:
                continue
            run = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            alignment = OxmlElement("w:vertAlign")
            alignment.set(qn("w:val"), "superscript")
            properties.append(alignment)
            run.append(properties)
            value = OxmlElement("w:t")
            value.text = "".join(marker.text or "" for marker in markers)
            run.append(value)
            math.addprevious(run)
            math.getparent().remove(math)


def main(input_path, output_path):
    document = Document(input_path)
    format_author_superscripts(document)
    capability = [table for table in document.tables
                  if len(table.columns) == 8
                  and table.cell(0, 0).text == "Software"]
    if len(capability) > 1:
        raise ValueError("Expected at most one related-software table")
    if capability:
        format_capability_table(capability[0])
    workflow = [table for table in document.tables
                if len(table.columns) == 6
                and table.cell(0, 0).text == "Method"
                and table.cell(0, 1).text == "Layout initialization"]
    if len(workflow) > 1:
        raise ValueError("Expected at most one workflow table")
    if workflow:
        format_workflow_table(workflow[0])
    seen_caption = False
    for paragraph in document.paragraphs:
        if not paragraph.text.startswith("Median elapsed time in seconds"):
            continue
        if seen_caption:
            paragraph._element.getparent().remove(paragraph._element)
        seen_caption = True
    wide = [table for table in document.tables if len(table.columns) == 12]
    if len(wide) not in (0, 3):
        raise ValueError("Expected zero or three wide runtime tables")
    if wide:
        template = document.element.body.sectPr
        add_section_break(wide[0]._element, template, landscape=False)
        after = OxmlElement("w:p")
        wide[-1]._element.addnext(after)
        add_section_break(after, template, landscape=True)
        after.getparent().remove(after)
        for table in wide:
            format_wide_table(table)
    document.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: format_jss_docx.py INPUT.docx OUTPUT.docx")
    main(sys.argv[1], sys.argv[2])
