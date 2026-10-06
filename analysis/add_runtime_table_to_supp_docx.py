#!/usr/bin/env python3

import csv
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def add_block(table, rows, columns, heading):
    row = table.add_row()
    row.cells[0].merge(row.cells[-1]).text = heading
    header = table.add_row()
    for cell, label in zip(header.cells, ("Method", "Scope", *columns)):
        cell.text = label
    for item in rows:
        row = table.add_row()
        method = item["method"].replace("fastEmbedR t-SNE ",
                                         "fastEmbedR ")
        method = method.replace("fastEmbedR PCA ", "fastEmbedR ")
        method = method.replace("fastEmbedR fuzzy UMAP ",
                                "fastEmbedR fuzzy ")
        method = method.replace("fastEmbedR binary UMAP ",
                                "fastEmbedR binary ")
        method = method.replace("RAPIDS cuML ", "cuML ")
        method = method.replace(" [Python]", "")
        for cell, value in zip(row.cells,
                               (method, item["scope"],
                                *(item[name] for name in columns))):
            cell.text = value


def main(source, csv_path, campaign_id, output):
    document = Document(source)
    anchor = next(paragraph for paragraph in document.paragraphs
                  if "combines the R and direct-Python results" in
                  paragraph.text)
    anchor.text = anchor.text.replace(
        "Table\u00a0[tab:all-method-performance]", "Table S1"
    )
    for paragraph in list(document.paragraphs):
        if paragraph.text.startswith("@p0.34lrrrrrr@") or \
                "Method & Scope & COIL-20" in paragraph.text:
            paragraph._element.getparent().remove(paragraph._element)

    caption = document.add_paragraph(
        "Table S1. Observed elapsed times (seconds), campaign "
        f"{campaign_id}. R and Python timing scopes differ."
    )
    caption.paragraph_format.page_break_before = True
    anchor._element.addnext(caption._element)
    table = document.add_table(rows=0, cols=8)
    table.style = "Table"
    caption._element.addnext(table._element)
    with open(csv_path, newline="", encoding="utf-8") as stream:
        rows = list(csv.DictReader(stream))
    datasets = list(rows[0])[3:]
    for family in ("pca", "tsne", "umap"):
        family_rows = [row for row in rows if row["family"] == family]
        for block, start in enumerate((0, 6), start=1):
            columns = datasets[start:start + 6]
            add_block(table, family_rows, columns,
                      f"{family.upper()}, data sets {block}/2")
    widths = [Inches(1.9), Inches(0.55)] + [Inches(0.6)] * 6
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = width
    for row in table.rows:
        properties = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        properties.append(no_split)
        for cell, width in zip(row.cells, widths):
            cell.width = width
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(7.5)
    for row in table.rows[:1]:
        properties = row._tr.get_or_add_trPr()
        repeated = OxmlElement("w:tblHeader")
        repeated.set(qn("w:val"), "true")
        properties.append(repeated)
    document.save(output)


if __name__ == "__main__":
    if len(sys.argv) != 5:
        raise SystemExit("Usage: add_runtime_table_to_supp_docx.py "
                         "SOURCE.docx RUNTIME.csv CAMPAIGN OUTPUT.docx")
    main(*sys.argv[1:])
